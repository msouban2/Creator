"""Build the Aaina admin dashboard presentation."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "images")
OUT = os.path.join(os.path.dirname(HERE), "Aaina-Admin-Dashboard.pptx")

# Brand palette
PLUM = RGBColor(0x7A, 0x1F, 0x3D)
ROSE = RGBColor(0xC2, 0x45, 0x6B)
CREAM = RGBColor(0xFB, 0xF4, 0xF3)
DARK = RGBColor(0x2B, 0x1B, 0x22)
GREY = RGBColor(0x6B, 0x5B, 0x62)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# slug -> (title, description)
SLIDES = [
    ("admin_login", "Admin Login", "Secure sign-in for administrators and employees."),
    ("admin_dashboard", "Dashboard", "At-a-glance KPIs: creators, active campaigns, applications and pending submissions."),
    ("admin_campaigns", "Campaigns", "Create and manage brand campaigns, budgets and status."),
    ("admin_applications", "Applications", "Review creator applications and approve or reject participation."),
    ("admin_users", "Users", "Directory of all creators with profiles and social reach."),
    ("admin_submissions", "Submissions", "Review submitted content and approve, request revisions or reject."),
    ("admin_kyc", "KYC Verification", "Verify creator identity documents before payouts."),
    ("admin_payments", "Payments", "Track and process creator payouts and payment history."),
    ("admin_referrals", "Referrals", "Monitor the referral program and commissions earned."),
    ("admin_notifications", "Notifications", "Broadcast and manage in-app notifications to creators."),
    ("admin_employees", "Employees", "Manage the admin team and add new team members."),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def fill(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, size, color,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Segoe UI"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return box


def add_bar(slide, left, top, width, height, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


# ---------- Title slide ----------
s = prs.slides.add_slide(blank)
fill(s, PLUM)
add_bar(s, 0, Inches(3.05), SW, Inches(0.06), ROSE)
add_text(s, Inches(1), Inches(2.0), Inches(11.33), Inches(1.2),
         "Aaina", 66, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.15), Inches(11.33), Inches(0.8),
         "Influencer Marketing Platform", 26, CREAM, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.1), Inches(11.33), Inches(0.6),
         "Admin Dashboard — Product Walkthrough", 18, ROSE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(6.6), Inches(11.33), Inches(0.5),
         "August 2026", 12, CREAM, align=PP_ALIGN.CENTER)

# ---------- Section divider ----------
s = prs.slides.add_slide(blank)
fill(s, CREAM)
add_bar(s, Inches(1), Inches(3.3), Inches(1.4), Inches(0.12), ROSE)
add_text(s, Inches(1), Inches(3.5), Inches(11), Inches(1.0),
         "Admin Dashboard", 44, PLUM, bold=True)
add_text(s, Inches(1.02), Inches(4.6), Inches(11), Inches(0.6),
         "The web console used by the Aaina team to run the marketplace.", 18, GREY)

# ---------- Content slides ----------
for slug, title, desc in SLIDES:
    path = os.path.join(IMG_DIR, slug + ".png")
    if not os.path.exists(path):
        continue
    s = prs.slides.add_slide(blank)
    fill(s, CREAM)
    # header
    add_bar(s, 0, 0, SW, Inches(1.1), PLUM)
    add_bar(s, 0, Inches(1.1), SW, Inches(0.05), ROSE)
    add_text(s, Inches(0.6), Inches(0.12), Inches(12), Inches(0.6),
             title, 28, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.62), Inches(0.66), Inches(12), Inches(0.4),
             desc, 13, CREAM, anchor=MSO_ANCHOR.MIDDLE)

    # image, fit within the area below the header, centered
    iw, ih = Image.open(path).size
    ratio = iw / ih
    max_w = SW - Inches(1.2)
    max_h = SH - Inches(1.7)
    w = max_w
    h = int(w / ratio)
    if h > max_h:
        h = max_h
        w = int(h * ratio)
    left = int((SW - w) / 2)
    top = Inches(1.1) + int((SH - Inches(1.1) - h) / 2)
    # subtle white card behind the image
    add_bar(s, left - Inches(0.08), top - Inches(0.08),
            w + Inches(0.16), h + Inches(0.16), WHITE)
    s.shapes.add_picture(path, left, top, width=w, height=h)

# ---------- Closing slide ----------
s = prs.slides.add_slide(blank)
fill(s, PLUM)
add_text(s, Inches(1), Inches(3.0), Inches(11.33), Inches(1.0),
         "Thank you", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.1), Inches(11.33), Inches(0.6),
         "Aaina — Influencer Marketing Platform", 18, CREAM, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("Saved:", OUT)
print("Slides:", len(prs.slides._sldIdLst))
